"""Show a buyer exactly where an extracted value came from: the cell, the page, the paragraph,
the email line, or the region of the photograph."""
import re
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw


def _attachment(msg: dict, kinds: tuple[str, ...]) -> dict | None:
    for a in msg.get("attachments", []):
        if a["name"].lower().endswith(kinds) or any(a["mime"].startswith(k) for k in kinds if "/" in k):
            return a
    return None


@lru_cache(maxsize=32)
def pdf_page_png(path: str, page: int) -> Image.Image | None:
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(path)
    if not 1 <= page <= len(doc):
        return None
    return doc[page - 1].render(scale=1.4).to_pil()


def image_region(path: str, region: list[float]) -> tuple[Image.Image, Image.Image]:
    """(full image with the region outlined, zoomed crop of the region)."""
    img = Image.open(path).convert("RGB")
    w, h = img.size
    x0, y0, x1, y1 = region
    pad = 0.02
    box = (max(0, (x0 - pad) * w), max(0, (y0 - pad) * h), min(w, (x1 + pad) * w), min(h, (y1 + pad) * h))
    crop = img.crop(tuple(int(v) for v in box))
    if crop.width < 500:
        f = 500 / max(crop.width, 1)
        crop = crop.resize((int(crop.width * f), int(crop.height * f)), Image.LANCZOS)
    full = img.copy()
    ImageDraw.Draw(full).rectangle([x0 * w, y0 * h, x1 * w, y1 * h], outline=(230, 30, 30), width=max(3, w // 250))
    full.thumbnail((700, 700))
    return full, crop


def xlsx_context(path: str, source: str) -> list[list[str]] | None:
    """The row(s) of the cited cell(s), as a small grid with the cited cell marked ▶."""
    m = re.findall(r"(?:'?([^'!]+)'?!)?\$?([A-Z]{1,3})\$?(\d+)", source)
    if not m:
        return None
    wb = load_workbook(path, data_only=False)
    sheet, col, row = m[0]
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    row = int(row)
    header_row = next((r for r in range(1, row) if sum(1 for c in ws[r] if c.value is not None) >= 4), None)
    out = []
    for r in ([header_row] if header_row else []) + [row]:
        out.append([("▶ " if (r == row and c.column_letter == col) else "") + ("" if c.value is None else str(c.value))
                    for c in ws[r]])
    return out


def docx_paragraph(path: str, source: str) -> str | None:
    from docx import Document
    m = re.search(r"paragraph\s*(\d+)", source, re.I)
    if not m:
        return None
    paras = [p.text for p in Document(path).paragraphs if p.text.strip()]
    i = int(m.group(1))
    return paras[i - 1] if 1 <= i <= len(paras) else None


def email_line(body: str, source: str) -> str | None:
    m = re.search(r"line\s*(\d+)", source, re.I)
    if not m:
        return None
    lines = body.splitlines()  # same numbering as extract._lines_text (blank lines count)
    i = int(m.group(1))
    return lines[i - 1] if 1 <= i <= len(lines) else None


def locate(msg: dict, field: dict) -> dict:
    """Work out what to show for a field. Returns {'kind': ..., ...}."""
    src = field.get("source") or ""
    low = src.lower()
    if field.get("region"):
        a = _attachment(msg, ("image/", ".jpg", ".jpeg", ".png", ".webp"))
        if a and Path(a["path"]).exists():
            return {"kind": "image", "path": a["path"], "region": field["region"]}
    if "email" in low and "line" in low:
        return {"kind": "text", "text": email_line(msg.get("body", ""), src)}
    m = re.search(r"page\s*(\d+)", low)
    if m:
        a = _attachment(msg, (".pdf",))
        if a and Path(a["path"]).exists():
            return {"kind": "pdf", "path": a["path"], "page": int(m.group(1))}
    if "!" in src or re.search(r"\b[A-Z]{1,2}\d{1,4}\b", src):
        a = _attachment(msg, (".xlsx", ".xlsm"))
        if a and Path(a["path"]).exists():
            return {"kind": "xlsx", "grid": xlsx_context(a["path"], src)}
    if "paragraph" in low:
        a = _attachment(msg, (".docx",))
        if a and Path(a["path"]).exists():
            return {"kind": "text", "text": docx_paragraph(a["path"], src)}
    if "photo" in low or "image" in low:
        a = _attachment(msg, ("image/", ".jpg", ".jpeg", ".png"))
        if a:
            return {"kind": "image_full", "path": a["path"]}
    return {"kind": "none"}
