"""Extraction: any vendor response in, structured JSON out, via one real model call.

One function for every vendor and every format. No per-vendor branches, no expected values.
The model decides what the document *says*: values, where each came from, how confident it
is, and which RFQ line a vendor item corresponds to (flagged as inferred when not literal).
It never converts currency, units, weeks→days or credit terms; that is normalise.py's job.

Rules enforced in Python after the call:
  - a value without a source is discarded and the field marked missing
  - nothing is filled in; missing is a visible state
  - confidence below REVIEW_THRESHOLD sends the field to the buyer's review queue
"""
import base64
import io
import json
import time
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image

from auth import get_secret
from llm import json_call, model_name
from rfq import standards_in

REVIEW_THRESHOLD = 0.8
# Thinking effort for extraction. Most of the latency is thinking; 'low' keeps accuracy on the
# answer-key eval (evals/extraction_eval.py) at a fraction of the time. Override with EXTRACT_EFFORT.
EXTRACT_EFFORT = "low"
MAX_IMAGE_PX = 2000


# ------------------------------------------------------------------ schema
def _field(desc: str) -> dict:
    return {
        "type": "object",
        "description": desc,
        "properties": {
            "value": {"type": "string",
                      "description": "Exactly what the document states, or \"\" if it does not state it. Numbers as "
                                     "plain decimals with '.' as decimal separator and no thousands separators "
                                     "(read '1.176,00' as 1176.00). Dates as YYYY-MM-DD."},
            "confidence": {"type": "number", "description": "0-1: how sure you are the value is what the document says"},
            "source": {"type": "string",
                       "description": "Where exactly: e.g. \"Offer!F8\", \"page 1, table row Pos. 20, column 'Unit price "
                                      "EUR'\", \"page 1, small print at foot\", \"paragraph 4\", \"email body, line 3\", "
                                      "\"photo, table row SL-ANU316, column 'Rate'\". \"\" only if value is \"\"."},
            "assumption": {"type": "string",
                           "description": "Any interpretation you made to read this value, else \"\"."},
            "region": {"type": "array", "items": {"type": "number"},
                       "description": "Images only: [x0, y0, x1, y1] as fractions 0-1 of image width/height "
                                      "bounding the text the value was read from. [] for non-image sources."},
        },
        "required": ["value", "confidence", "source", "assumption", "region"],
        "additionalProperties": False,
    }


LINE = {
    "type": "object",
    "properties": {
        "rfq_line": {"type": "integer",
                     "description": "Index of the RFQ line this vendor item answers, or -1 if it answers none."},
        "vendor_item": {"type": "string", "description": "Vendor's own part number / description, verbatim."},
        "mapping": {"type": "string", "enum": ["stated", "inferred"],
                    "description": "stated = the document names the RFQ variant/material; inferred = you matched it "
                                   "from a description, series or part number."},
        "mapping_note": {"type": "string", "description": "Why this item maps to that RFQ line."},
        "quoted": {"type": "boolean", "description": "false if the vendor explicitly declines this line."},
        "unit_price": _field("Price per vendor unit, as stated."),
        "currency": _field("ISO code of the price currency (INR for Rs/₹), as stated or clearly implied."),
        "units_per_price": _field("How many RFQ units (e.g. single bearing assemblies) one priced vendor unit "
                                  "contains, as stated — e.g. '2' for 'SET (2 PCS)'. '1' when priced per piece."),
        "qty_offered": _field("Quantity offered, in the vendor's units, as stated."),
        "availability": _field("One of: full, partial, not_available — as stated or directly implied."),
        "lead_time": _field("Lead time text as stated, e.g. '5 weeks', '30 days', '12-13 weeks'."),
        "lead_time_basis": _field("One of: delivered (arrives at buyer), ex_works (ready at vendor), dispatch."),
        "balance": _field("If partial: what the vendor says about the remaining quantity, verbatim, else null."),
        "remarks": {"type": "string"},
    },
    "required": ["rfq_line", "vendor_item", "mapping", "mapping_note", "quoted", "unit_price", "currency",
                 "units_per_price", "qty_offered", "availability", "lead_time", "lead_time_basis", "balance",
                 "remarks"],
    "additionalProperties": False,
}

CREDIT_KINDS = ["net_days", "discount_then_net", "advance_split", "lc_at_sight", "days_from_grn",
                "refers_to_history", "other", "unstated"]
FREIGHT_KINDS = ["delivered", "ex_works", "conditional", "unstated"]

SCHEMA = {
    "type": "object",
    "properties": {
        "summary": {"type": "string", "description": "Two sentences: what this response is and anything unusual."},
        "vendor_quote_ref": _field("The vendor's own quotation / offer number."),
        "lines": {"type": "array", "items": LINE,
                  "description": "One entry per vendor item that answers an RFQ line, plus one entry with "
                                 "quoted=false for each RFQ line the vendor explicitly declines. Ignore catalogue "
                                 "items that answer no RFQ line."},
        "credit": {
            "type": "object",
            "properties": {
                "text": _field("Payment / credit terms, verbatim."),
                "kind": {"type": "string", "enum": CREDIT_KINDS},
                "net_days": _field("Net credit days as stated, e.g. '45', '30'."),
                "discount_pct": _field("Early-payment discount %, e.g. '2'."),
                "discount_days": _field("Days within which the discount applies, e.g. '10'."),
                "advance_pct": _field("Advance payment %, e.g. '50'."),
                "balance_event": _field("When the balance / net period runs from: invoice, delivery, grn, shipment."),
            },
            "required": ["text", "kind", "net_days", "discount_pct", "discount_days", "advance_pct", "balance_event"],
            "additionalProperties": False,
        },
        "freight": {
            "type": "object",
            "properties": {
                "text": _field("Freight / price basis terms, verbatim."),
                "kind": {"type": "string", "enum": FREIGHT_KINDS},
                "threshold_amount": _field("For conditional freight: the order/despatch value threshold as a plain "
                                           "number (e.g. '500000' for '₹5 lakh')."),
                "threshold_currency": _field("Currency of that threshold."),
            },
            "required": ["text", "kind", "threshold_amount", "threshold_currency"],
            "additionalProperties": False,
        },
        "volume_discount": {
            "type": "object",
            "properties": {"text": _field("Any volume rebate / discount clause, verbatim."),
                           "pct": _field("Percentage."), "threshold_amount": _field("Order value threshold."),
                           "threshold_currency": _field("Currency of the threshold.")},
            "required": ["text", "pct", "threshold_amount", "threshold_currency"],
            "additionalProperties": False,
        },
        "validity": _field("Quote validity as stated."),
        "cert_assessment": {
            "type": "array",
            "description": "One entry per REQUIRED standard listed in the RFQ context.",
            "items": {
                "type": "object",
                "properties": {
                    "required": {"type": "string", "description": "The required standard, as given in the context."},
                    "material": {"type": "string"},
                    "match": {"type": "string", "enum": ["exact", "near_equivalent", "different", "not_provided"],
                              "description": "exact = the vendor certifies to this very standard; near_equivalent = "
                                             "a related but different standard; different = unrelated; "
                                             "not_provided = nothing in the document covers it."},
                    "vendor_standard": _field("The standard the vendor actually cites."),
                    "certificate_no": _field("Certificate / report number."),
                    "issuer": _field("Issuing body."),
                    "expiry": _field("Expiry / valid-until date, YYYY-MM-DD."),
                    "note": {"type": "string", "description": "Why you judged the match this way."},
                },
                "required": ["required", "material", "match", "vendor_standard", "certificate_no", "issuer", "expiry",
                             "note"],
                "additionalProperties": False,
            },
        },
        "history_references": {
            "type": "array",
            "description": "Places where the vendor refers to past dealings instead of stating terms, e.g. 'same "
                           "terms as our last order'. Do not resolve them.",
            "items": {"type": "object",
                      "properties": {"quote": {"type": "string"}, "source": {"type": "string"},
                                     "affects": {"type": "array", "items": {"type": "string",
                                                 "enum": ["credit", "freight", "price", "delivery", "certification",
                                                          "other"]}}},
                      "required": ["quote", "source", "affects"], "additionalProperties": False},
        },
        "ambiguities": {"type": "array", "items": {"type": "string"},
                        "description": "Anything a careful buyer would want to query."},
    },
    "required": ["summary", "vendor_quote_ref", "lines", "credit", "freight", "volume_discount", "validity",
                 "cert_assessment", "history_references", "ambiguities"],
    "additionalProperties": False,
}


# ------------------------------------------------------------------ document → model content
def _xlsx_text(p: Path) -> str:
    wb = load_workbook(p, data_only=False)
    cached = load_workbook(p, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== Sheet '{ws.title}' (cell-by-cell; empty cells omitted) ===")
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            out.append("Merged ranges: " + ", ".join(merged))
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    cv = cached[ws.title][c.coordinate].value
                    v = f"{v} (formula; cached value: {cv if cv is not None else 'none saved'})"
                cells.append(f"{c.coordinate}={v!r}")
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _docx_text(p: Path) -> str:
    d = Document(p)
    out = [f"paragraph {i}: {para.text}" for i, para in enumerate(d.paragraphs, 1) if para.text.strip()]
    for ti, t in enumerate(d.tables, 1):
        for ri, row in enumerate(t.rows, 1):
            out.append(f"table {ti} row {ri}: " + " | ".join(c.text for c in row.cells))
    return "\n".join(out)


def _lines_text(text: str, label: str) -> str:
    return "\n".join(f"{label}, line {i}: {l}" for i, l in enumerate(text.splitlines(), 1) if l.strip())


def _image_block(p: Path) -> dict:
    img = Image.open(p)
    img = img.convert("RGB")
    if max(img.size) > MAX_IMAGE_PX:
        img.thumbnail((MAX_IMAGE_PX, MAX_IMAGE_PX))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90)
    return {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                        "data": base64.standard_b64encode(buf.getvalue()).decode()}}


def document_blocks(message: dict) -> tuple[list[dict], list[str]]:
    """Every part of an inbound message as model content blocks. Returns (blocks, inventory)."""
    blocks, inventory = [], []
    body = (message.get("body") or "").strip()
    if body:
        blocks.append({"type": "text", "text": f"--- EMAIL (from {message['from_addr']}, subject "
                                               f"'{message['subject']}') ---\n{_lines_text(body, 'email body')}"})
        inventory.append("email body")
    for a in message.get("attachments", []):
        p, name, mime = Path(a["path"]), a["name"], a["mime"]
        low = name.lower()
        if low.endswith(".pdf") or mime == "application/pdf":
            blocks.append({"type": "text", "text": f"--- ATTACHMENT '{name}' (PDF, cite page numbers) ---"})
            blocks.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                          "data": base64.standard_b64encode(p.read_bytes()).decode()}})
        elif mime.startswith("image/") or low.endswith((".jpg", ".jpeg", ".png", ".webp")):
            blocks.append({"type": "text", "text": f"--- ATTACHMENT '{name}' (photograph / image; give a region for "
                                                   "every value) ---"})
            blocks.append(_image_block(p))
        elif low.endswith((".xlsx", ".xlsm")):
            blocks.append({"type": "text", "text": f"--- ATTACHMENT '{name}' (spreadsheet; cite cells as "
                                                   f"Sheet!A1) ---\n{_xlsx_text(p)}"})
        elif low.endswith(".docx"):
            blocks.append({"type": "text", "text": f"--- ATTACHMENT '{name}' (Word document; cite paragraph "
                                                   f"numbers) ---\n{_docx_text(p)}"})
        elif low.endswith((".txt", ".csv", ".eml")) or mime.startswith("text/"):
            blocks.append({"type": "text", "text": f"--- ATTACHMENT '{name}' (text) ---\n"
                                                   f"{_lines_text(p.read_text(errors='replace'), name)}"})
        else:
            inventory.append(f"{name} (unsupported format, not read)")
            continue
        inventory.append(name)
    return blocks, inventory


# ------------------------------------------------------------------ the call
def _context(rfq: dict, vendor: dict, received: str) -> str:
    lines = "\n".join(f"  [{i}] {l['material']} — variant {l['variant']} — qty {l['qty']} {l['uom']} — need-by "
                      f"{l['need_by']}" for i, l in enumerate(rfq["lines"]))
    reqs = "\n".join(f"  - {m}: {p.strip()}" for m, q in standards_in(rfq) for p in q.split("+") if p.strip())
    return f"""RFQ {rfq['rfq_id']} ({rfq['title']}) from {rfq['buyer']['company']}.
Price basis requested: {rfq['terms']['price_basis']}.
Credit terms basis requested: {rfq['terms']['credit_terms']}.
RFQ lines (use these indexes for rfq_line):
{lines}
Required quality standards — one cert_assessment entry per line below, judged separately (a material with two
requirements gets two entries; "exact" only if that specific requirement is evidenced):
{reqs}
The response below was received from vendor {vendor['code']} ({vendor['name']}) on {received}."""


SYSTEM = """You read vendor quotation documents for an industrial procurement team and record what they say.

You are a careful reader, not a calculator:
- Record values exactly as the document states them. Do not convert currencies, units, weeks to days, or credit terms
  to numbers of days. Do not multiply, divide or total anything. Python does all arithmetic afterwards.
- Every value needs a source precise enough for a buyer to find it (sheet+cell, page+row+column, paragraph, email
  line, photo row). If you cannot point to where a value came from, return value "".
- Never fill a gap with a plausible value. If the document does not state something, value is "" (empty string).
  Missing is fine.
- Confidence reflects legibility and ambiguity: clean printed/tabular value ~0.95+; value read from small print,
  prose or a photo ~0.8-0.9; value that needed interpretation ~0.6-0.8; guesswork is not allowed at all.
- Map each vendor item to an RFQ line. If the document names the variant, mapping=stated. If you matched on a series
  number, description ('general duty', 'HD type', 'Schwerlast') or part number, mapping=inferred and say why.
- If a vendor explicitly declines a line, include it with quoted=false. If a vendor is silent about a line, omit it.
- For certifications, compare what the vendor cites against each required standard. A related but different
  standard (e.g. DIN vs ISO bearing standards) is near_equivalent, never exact. A product test report does not
  cover a supplier quality-system requirement such as ISO 9001.
- If the vendor refers to past dealings instead of stating terms ('same as last order'), record it in
  history_references and leave the affected fields "". Do not resolve it.
- Read small print, footers and merged cells: terms are often buried there.
- A lead time, price basis or other term stated once for the whole document (a footer, a covering note, a
  "Delivery:" line) applies to every item it covers: record it on each of those items, citing that one source.
  If it explicitly covers only some item types, apply it only to those."""


def extract(message: dict, rfq: dict, vendor: dict, today: date | None = None) -> dict:
    """The one extraction function. Returns {'raw': model output, 'meta': {...}, 'discarded': [...]}."""
    received = message.get("sent_date") or (today or date.today()).isoformat()
    blocks, inventory = document_blocks(message)
    if not blocks:
        raise ValueError("Nothing readable in this message")
    content = [{"type": "text", "text": _context(rfq, vendor, received)}] + blocks + [
        {"type": "text", "text": "Extract this response into the schema. Remember: sources for everything, "
                                 "no arithmetic, \"\" when not stated."}]
    t0 = time.time()
    raw = json_call(SYSTEM, [{"role": "user", "content": content}], SCHEMA, effort=get_secret("EXTRACT_EFFORT") or EXTRACT_EFFORT,
                    max_tokens=32000, constrained=False,
                    )
    _to_nulls(raw)
    discarded = enforce_sources(raw)
    return {"raw": raw, "discarded": discarded,
            "meta": {"model": model_name(), "seconds": round(time.time() - t0, 1), "documents": inventory,
                     "extracted_at": time.strftime("%Y-%m-%d %H:%M:%S")}}


def _walk_fields(obj, path=""):
    """Yield (path, field_dict) for every value/confidence/source field in the extraction."""
    if isinstance(obj, dict):
        if {"value", "confidence", "source"} <= set(obj):
            yield path, obj
            return
        for k, v in obj.items():
            yield from _walk_fields(v, f"{path}.{k}" if path else k)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from _walk_fields(v, f"{path}[{i}]")


def _to_nulls(raw: dict) -> None:
    """The schema uses "" / [] / -1 for 'not stated' (the API limits nullable fields); restore real nulls."""
    for _, f in _walk_fields(raw):
        for k in ("value", "source", "assumption"):
            if isinstance(f.get(k), str) and not f[k].strip():
                f[k] = None
        if not f.get("region") or len(f["region"]) != 4:
            f["region"] = None
    for l in raw.get("lines", []):
        if l.get("rfq_line", -1) < 0:
            l["rfq_line"] = None


def enforce_sources(raw: dict) -> list[dict]:
    """A value with no source is discarded, not trusted. Mutates raw; returns what was dropped."""
    dropped = []
    for path, f in _walk_fields(raw):
        if f.get("value") not in (None, "") and not (f.get("source") or "").strip():
            dropped.append({"field": path, "value": f["value"]})
            f["value"] = None
            f["assumption"] = "Discarded: the model gave no source for this value"
    return dropped


def review_items(raw: dict) -> list[dict]:
    """Fields a buyer must confirm: stated but below the confidence threshold."""
    return [{"field": p, **f} for p, f in _walk_fields(raw)
            if f.get("value") not in (None, "") and float(f.get("confidence") or 0) < REVIEW_THRESHOLD]


def get_field(raw: dict, path: str):
    cur = raw
    for part in path.replace("]", "").replace("[", ".").split("."):
        cur = cur[int(part)] if part.isdigit() else cur[part]
    return cur


def dumps(x) -> str:
    return json.dumps(x, ensure_ascii=False)
