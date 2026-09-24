"""Render an RFQ JSON object into the xlsx pack sent to vendors. Plain openpyxl, no AI."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from master_data import MATERIALS
from rfq import materials_in, standards_in

NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # shaded = vendor fills in
LABEL_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def _title(ws, rfq: dict, subtitle: str) -> int:
    b = rfq["buyer"]
    ws["A1"] = b["company"]
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = b["plant"]
    ws["A3"] = f"Request for Quotation {rfq['rfq_id']} — {rfq['title']}"
    ws["A3"].font = Font(bold=True, size=12)
    ws["A4"] = subtitle
    ws["A4"].font = Font(italic=True, color="595959")
    return 6


def _header_row(ws, row: int, headers: list[str], input_from: int | None = None) -> None:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = WHITE_BOLD
        c.fill = HEADER_FILL if input_from is None or i < input_from else PatternFill("solid", fgColor="BF8F00")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX


def _terms_sheet(wb: Workbook, rfq: dict) -> None:
    ws = wb.active
    ws.title = "Terms"
    r = _title(ws, rfq, "Sheet 1 of 3 — terms and instructions. Please complete sheets 2 and 3.")
    b = rfq["buyer"]
    t = rfq["terms"]
    mats = standards_in(rfq)
    rows = [
        ("RFQ reference", rfq["rfq_id"]),
        ("Scope", f"{rfq['title']} — {len(rfq['lines'])} item(s), see sheet 2"),
        ("Quality standards required", "\n".join(f"{m}: {q}" for m, q in mats)),
        ("Specification notes", "\n".join(f"{MATERIALS[c].name}: {MATERIALS[c].spec_notes}"
                                           for c in materials_in(rfq) if MATERIALS[c].spec_notes) or "—"),
        ("Planned PO date", rfq.get("planned_po_date", "—")),
        ("Quotes close", rfq["close_date"]),
        ("Deliver to", b["plant"]),
        ("Contact", f"{b['contact']} · {b['email']} · {b['phone']}"),
        (None, None),
        ("RESPONSE BASIS", "Please quote on these bases so responses can be compared like for like"),
        ("Unit price", t["price_basis"]),
        ("Currency", t["currency"]),
        ("Freight", t["freight"]),
        ("Delivery time", t["delivery"]),
        ("Credit terms (DSO)", t["credit_terms"]),
        ("Availability", t["availability"]),
        ("Quality certification", t["quality"]),
        ("Validity", t["validity"]),
    ]
    if rfq.get("notes"):
        rows.append(("Buyer notes", rfq["notes"]))
    rows += [
        (None, None),
        ("How to respond", "Reply to the RFQ email with this workbook completed. Shaded cells are for you to fill. "
                           "If you prefer your own format, that is fine — please cover every item above."),
    ]
    for label, value in rows:
        if label is not None:
            a = ws.cell(row=r, column=1, value=label)
            a.font = Font(bold=True)
            a.fill = LABEL_FILL
            a.border = BOX
            a.alignment = WRAP
            v = ws.cell(row=r, column=2, value=value)
            v.alignment = WRAP
            v.border = BOX
            if label == "RESPONSE BASIS":
                a.fill = HEADER_FILL
                a.font = WHITE_BOLD
        r += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 95


LINE_HEADERS = ["Item", "Material", "Variant", "Qty required", "UoM", "Need-by date",
                "Your part no.", "Unit price", "Currency", "Price is per (unit basis)",
                "Freight included?", "Qty you can supply", "Availability",
                "Delivery (days from PO)", "Remarks"]
LINE_INPUT_FROM = 7  # 1-indexed column where vendor inputs start


def _lines_sheet(wb: Workbook, rfq: dict) -> None:
    ws = wb.create_sheet("Items")
    r = _title(ws, rfq, "Sheet 2 of 3 — one row per item. Fill the shaded columns.")
    _header_row(ws, r, LINE_HEADERS, input_from=LINE_INPUT_FROM)
    ws.row_dimensions[r].height = 32
    first = r + 1
    for i, v in enumerate(rfq["lines"], start=1):
        row = r + i
        vals = [i, v["material"], v["variant"], v["qty"], v.get("uom", "nos"), date.fromisoformat(v["need_by"])]
        for col, val in enumerate(vals, start=1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = BOX
        ws.cell(row=row, column=6).number_format = "DD-MMM-YYYY"
        for col in range(LINE_INPUT_FROM, len(LINE_HEADERS) + 1):
            c = ws.cell(row=row, column=col)
            c.fill = INPUT_FILL
            c.border = BOX
        ws.cell(row=row, column=8).number_format = "#,##0.00"
    last = r + len(rfq["lines"])

    def dv(options: list[str], col_letter: str) -> None:
        d = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(d)
        d.add(f"{col_letter}{first}:{col_letter}{last}")

    dv(["INR", "USD", "EUR", "GBP", "JPY", "CNY"], "I")
    dv(["Yes - delivered", "No - ex-works", "Conditional"], "K")
    dv(["Full", "Partial", "Not available"], "M")

    # Commercial terms that apply to the whole quote
    r = last + 2
    ws.cell(row=r, column=1, value="Commercial terms (whole quote)").font = Font(bold=True, color=NAVY)
    for label in ["Credit period, days from invoice date",
                  "Early-payment discount (if any)",
                  "Freight conditions (if conditional)",
                  "Quote valid until"]:
        r += 1
        a = ws.cell(row=r, column=1, value=label)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        a.fill = LABEL_FILL
        a.font = Font(bold=True)
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = BOX
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=11)
        for col in range(7, 12):
            ws.cell(row=r, column=col).fill = INPUT_FILL
            ws.cell(row=r, column=col).border = BOX

    widths = [6, 30, 18, 12, 7, 13, 16, 12, 10, 20, 17, 14, 14, 14, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = ws.cell(row=first, column=4)


CERT_HEADERS = ["Material", "Requirement", "Certified to (standard + edition)", "Certificate no.",
                "Issuing body", "Valid until", "Copy attached?"]


def _certs_sheet(wb: Workbook, rfq: dict) -> None:
    ws = wb.create_sheet("Certifications")
    r = _title(ws, rfq, "Sheet 3 of 3 — one row per required standard. If you certify to a different "
                        "or equivalent standard, say so here rather than leaving it blank.")
    _header_row(ws, r, CERT_HEADERS, input_from=3)
    parts = [(m, p.strip()) for m, q in standards_in(rfq) for p in q.split("+") if p.strip()]
    for i, (mat, p) in enumerate(parts, start=1):
        row = r + i
        for col, val in ((1, mat), (2, p)):
            c = ws.cell(row=row, column=col, value=val)
            c.border = BOX
            c.fill = LABEL_FILL
        for col in range(3, len(CERT_HEADERS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.border = BOX
        ws.cell(row=row, column=6).number_format = "DD-MMM-YYYY"
    d = DataValidation(type="list", formula1='"Yes,No,To follow"', allow_blank=True)
    ws.add_data_validation(d)
    d.add(f"G{r + 1}:G{r + len(parts)}")
    note = r + len(parts) + 2
    ws.cell(row=note, column=1, value="Declaration: we confirm the certificates above are current and "
                                       "cover the items quoted.").font = Font(italic=True)
    ws.cell(row=note + 2, column=1, value="Name / designation:").font = Font(bold=True)
    ws.cell(row=note + 3, column=1, value="Date:").font = Font(bold=True)
    for rr in (note + 2, note + 3):
        ws.cell(row=rr, column=2).fill = INPUT_FILL
        ws.cell(row=rr, column=2).border = BOX
    for col, w in zip("ABCDEFG", [32, 34, 32, 20, 24, 14, 14]):
        ws.column_dimensions[col].width = w


def render_rfq_xlsx(rfq: dict) -> bytes:
    wb = Workbook()
    _terms_sheet(wb, rfq)
    _lines_sheet(wb, rfq)
    _certs_sheet(wb, rfq)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
