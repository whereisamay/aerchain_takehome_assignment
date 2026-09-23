"""Fabricate the six vendor responses to RFQ-2026-MCH-005 (bearing assembly) as real files.

This script and DATASET_NOTES.md live outside app/ on purpose. The application never imports
or reads anything in dataset/ except the vendor response files themselves, which arrive through
the (simulated) mail inbox exactly like real vendor replies.

Run:  python dataset/generate_vendor_files.py
"""
from pathlib import Path

import numpy as np
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageEnhance, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle

OUT = Path(__file__).resolve().parent / "vendor_responses"
OUT.mkdir(exist_ok=True)

pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("DejaVuSerif", "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"))

THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# --------------------------------------------------------------------------- Vendor A
def vendor_a() -> Path:
    """xlsx, vendor's own column order and part numbers. Partial qty on high-speed."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Offer"
    ws["A1"] = "SHREEJI BEARINGS & TRANSMISSION PVT LTD"
    ws["A1"].font = Font(bold=True, size=14, color="7B2C2C")
    ws["A2"] = "Survey No. 118, Metoda GIDC, Rajkot 360021 · GSTIN 24AAECS0000B1Z2"
    ws["A4"] = "Offer No: SBT/OF/26-27/0642"
    ws["E4"] = "Date: 28-09-2026"
    ws["A5"] = "Kind Attn: Procurement Desk, Deccan Flow Systems Pvt Ltd, Chakan"
    ws["E5"] = "Your Ref: RFQ-2026-MCH-005"
    headers = ["S.No", "Item Code", "Description", "Offered Qty", "UOM", "Rate (INR)", "Disc %", "Net Rate",
               "Lead Time", "Remarks"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=7, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7B2C2C")
        c.border = BOX
    rows = [
        [1, "SBT-NU316-HD", "Cyl. roller brg assy, heavy duty (NU316 ECJ/C3), with housing", 24, "Nos",
         4820, 0, 4820, "5 weeks", ""],
        [2, "SBT-6212-2RS", "Deep groove ball brg assy, general duty (6212-2RS/C3), with housing", 40, "Nos",
         2450, 0, 2450, "5 weeks", ""],
        [3, "SBT-7014-HS", "Angular contact brg assy, high speed (7014 ACD/P4A)", 12, "Nos",
         6700, 0, 6700, "5 weeks", "Can offer 60% of reqd qty (12 of 20) in 5 wks. Balance 8 nos in 11 wks."],
    ]
    for r, row in enumerate(rows, start=8):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (6, 8):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    t = 12
    terms = [
        ("TERMS & CONDITIONS", None),
        ("Price basis", "FOR your Chakan plant — freight & transit insurance included. GST @18% extra."),
        ("Payment", "45 days net"),
        ("Validity", "60 days"),
        ("Quality", "Bearings tested & classified per ISO 15243. Test cert no. SBT/QC/15243/2026-118, "
                    "valid to 31-Mar-2028."),
        ("QMS", "ISO 9001:2015, TUV Nord cert no. IN-QMS-44871, valid to 14-Jan-2028."),
        ("Note", "Item codes above are our catalogue codes. Housing included in all assemblies."),
    ]
    for label, val in terms:
        ws.cell(row=t, column=2, value=label).font = Font(bold=True)
        if val:
            ws.cell(row=t, column=3, value=val)
        t += 1
    ws.cell(row=t + 1, column=2, value="For Shreeji Bearings & Transmission Pvt Ltd — Authorised Signatory")
    for col, w in zip("ABCDEFGHIJ", [6, 16, 58, 12, 7, 12, 8, 12, 11, 44]):
        ws.column_dimensions[col].width = w
    p = OUT / "vendor_A_Shreeji_offer.xlsx"
    wb.save(p)
    return p


# --------------------------------------------------------------------------- Vendor B
def vendor_b() -> Path:
    """PDF on letterhead, EUR, near-equivalent standard (DIN 628), terms buried in small print."""
    p = OUT / "vendor_B_Nordlager_Angebot.pdf"
    doc = SimpleDocTemplate(str(p), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm)
    s = lambda size, font="DejaVu", color=colors.black, lead=None: ParagraphStyle(  # noqa: E731
        "x", fontName=font, fontSize=size, textColor=color, leading=lead or size * 1.3)
    story = [
        Paragraph("NORDLAGER", s(26, "DejaVu-Bold", colors.HexColor("#0B4F8A"))),
        Paragraph("Antriebstechnik GmbH · Wälzlager und Lagereinheiten seit 1962", s(9, color=colors.grey)),
        Paragraph("Heilbronner Straße 211 · 70191 Stuttgart · Germany · export@nordlager.example · "
                  "USt-IdNr. DE000000000", s(8, color=colors.grey)),
        Spacer(1, 10 * mm),
        Paragraph("Deccan Flow Systems Pvt Ltd<br/>Procurement Desk<br/>Plot 42, MIDC Chakan Phase II<br/>"
                  "Pune 410501, India", s(10)),
        Spacer(1, 6 * mm),
        Paragraph("<b>Angebot / Quotation NL-Q-2026-3318</b>", s(13, "DejaVu-Bold")),
        Paragraph("Stuttgart, 29.09.2026 · Your enquiry RFQ-2026-MCH-005 dated 23.09.2026", s(9)),
        Spacer(1, 5 * mm),
        Paragraph("Dear Sir or Madam,<br/>thank you for your enquiry. We are pleased to offer as follows:", s(10)),
        Spacer(1, 4 * mm),
    ]
    data = [["Pos.", "Article", "Description", "Qty", "Unit price\nEUR", "Total\nEUR"],
            ["10", "NL-LE 6212-G", "Lagereinheit Standard, Rillenkugellager 6212", "40", "24,50", "980,00"],
            ["20", "NL-LE NU316-S", "Lagereinheit Schwerlast, Zylinderrollenlager NU316", "24", "49,00", "1.176,00"],
            ["30", "NL-LE 7014-HG", "Lagereinheit Hochgeschwindigkeit, Schrägkugellager 7014", "20", "66,00",
             "1.320,00"],
            ["", "", "", "", "Summe", "3.476,00"]]
    cell = s(8.5)
    data = [[Paragraph(v, cell) if i == 2 and r > 0 else v for i, v in enumerate(row)]
            for r, row in enumerate(data)]
    tbl = Table(data, colWidths=[11 * mm, 27 * mm, 76 * mm, 11 * mm, 21 * mm, 22 * mm])
    tbl.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "DejaVu", 8.5),
        ("FONT", (0, 0), (-1, 0), "DejaVu-Bold", 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F2")),
        ("GRID", (0, 0), (-1, -2), 0.4, colors.grey),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("FONT", (4, -1), (-1, -1), "DejaVu-Bold", 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story += [tbl, Spacer(1, 6 * mm)]
    story += [
        Paragraph("<b>Delivery:</b> 6 weeks after receipt of order, ex works.", s(10)),
        Paragraph("<b>Quality:</b> All bearing units are manufactured and tested in accordance with DIN 628 "
                  "and DIN ISO 281. Our company holds ISO 9001:2015 certification (DQS certificate "
                  "074512 QM15, valid until 31.05.2028).", s(10)),
        Spacer(1, 4 * mm),
        Paragraph("We would be delighted to receive your order and remain at your disposal for any questions.",
                  s(10)),
        Spacer(1, 4 * mm),
        Paragraph("Mit freundlichen Grüßen / Kind regards<br/><b>Jürgen Albrecht</b><br/>Export Sales Asia",
                  s(10)),
        Spacer(1, 40 * mm),
        Paragraph(
            "Allgemeine Lieferbedingungen: Es gelten ausschließlich unsere AGB. Terms of payment: 2% discount if "
            "paid within 10 days of invoice date, otherwise net 30 days from invoice date. Prices EXW Stuttgart "
            "(Incoterms® 2020), incl. standard export packing; freight, insurance, customs clearance and import "
            "duties for buyer's account. A volume rebate of 3% applies to single orders exceeding EUR 5.000 net. "
            "Offer valid 45 days. Retention of title until full payment. Place of jurisdiction Stuttgart. "
            "Registergericht Stuttgart HRB 000000 · Geschäftsführer: Dr. K. Weiß.",
            s(5.5, color=colors.HexColor("#808080"), lead=7)),
    ]
    doc.build(story)
    return p


# --------------------------------------------------------------------------- Vendor C
def vendor_c() -> Path:
    """docx letter, commercials in prose. 2 of 3 variants. Cert expires before need-by.
    50% advance terms. Conditional freight."""
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    h = d.add_paragraph()
    r = h.add_run("KAVERI PRECISION COMPONENTS")
    r.bold = True
    r.font.size = Pt(16)
    d.add_paragraph("SF No. 44/2, Avinashi Road, Peelamedu, Coimbatore 641004 · info@kaveriprecision.example")
    d.add_paragraph("Ref: KPC/MKT/2026/1187                                                            "
                    "30 September 2026")
    d.add_paragraph("To\nThe Procurement Manager\nDeccan Flow Systems Pvt Ltd\nChakan, Pune")
    d.add_paragraph("Sub: Your enquiry for bearing assemblies — RFQ-2026-MCH-005").runs[0].bold = True
    paras = [
        "Dear Sir,",
        "We thank you for your valued enquiry and for considering Kaveri Precision Components as a potential "
        "supplier. Having studied your requirement in detail, we are pleased to submit our most competitive "
        "offer as below.",
        "For the standard bearing assembly, we are able to offer a price of Rupees Two Thousand Three Hundred "
        "and Eighty only (₹2,380) per assembly, for the full quantity of 40 numbers. For the heavy-duty bearing "
        "assembly, our price is ₹4,650 per assembly, again for the full quantity of 24 numbers required by you.",
        "We regret to inform you that we do not presently manufacture the high-speed variant and are therefore "
        "unable to quote for the same at this time. We hope to add this range to our portfolio next year.",
        "Delivery of both items can be effected within 30 days from the date of your purchase order. As regards "
        "commercial terms, we request 50% advance along with the purchase order, with the balance payable on "
        "delivery of material. Freight will be free to your Chakan plant for any single despatch above ₹5 lakh "
        "in value; for despatches below this value, freight will be charged at actuals. GST at the applicable "
        "rate will be extra.",
        "Our bearing assemblies conform to ISO 15243, certificate number KPC/ISO15243/0457 issued by Bureau "
        "Veritas, valid until 12 November 2026. We are also an ISO 9001:2015 certified organisation "
        "(certificate 9001-IN-22817, valid till 31 August 2027). Copies of both certificates can be shared on "
        "request.",
        "This offer is valid for 30 days. We look forward to the privilege of serving you.",
        "Thanking you,\nYours faithfully,\nfor Kaveri Precision Components\n\nS. Muthukumar\nManager — Marketing",
    ]
    for t in paras:
        d.add_paragraph(t)
    p = OUT / "vendor_C_Kaveri_letter.docx"
    d.save(p)
    return p


# --------------------------------------------------------------------------- Vendor D
def vendor_d() -> Path:
    """xlsx with junk rows, merged cells, header at row 7. USD, priced per set of 2.
    LC at sight. No freight stated. Product test report but no ISO 9001."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A1:H1")
    ws["A1"] = "PACIFIC MOTION COMPONENTS PTE LTD"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:H2")
    ws["A2"] = "21 Tuas South Ave 3, Singapore 637419 | UEN 201900000K | quotes@pacificmotion.example"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A3:H3")
    ws["A3"] = "Q U O T A T I O N"
    ws["A3"].font = Font(bold=True, size=13)
    ws["A3"].alignment = Alignment(horizontal="center")
    ws["A4"] = "Quote #: PMC-2609-0417"
    ws["F4"] = "Date: 29/09/2026"
    ws.merge_cells("A5:H5")
    ws["A5"] = "To: Deccan Flow Systems Pvt Ltd (India)   Attn: Procurement   Re: bearing assemblies"
    ws["A6"] = None
    headers = ["Item", "Part No.", "Desc.", "UOM", "Qty", "Unit Price (USD)", "Amount (USD)", "Lead Time"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=7, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
        c.border = BOX
    rows = [
        [1, "PMC-BA-6212", "Bearing assy - std type", "SET (2 PCS)", 20, 54.00],
        [2, "PMC-BA-NU316", "Bearing assy - HD type", "SET (2 PCS)", 12, 108.00],
        [3, "PMC-BA-7014H", "Bearing assy - HS type", "SET (2 PCS)", 10, 148.00],
    ]
    for r, row in enumerate(rows, start=8):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v).border = BOX
        ws.cell(row=r, column=7, value=f"=E{r}*F{r}").border = BOX
        ws.cell(row=r, column=8, value="4 wks").border = BOX
        for c in (6, 7):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    ws.merge_cells("A11:F11")
    ws["A11"] = "TOTAL"
    ws["A11"].alignment = Alignment(horizontal="right")
    ws["A11"].font = Font(bold=True)
    ws["G11"] = "=SUM(G8:G10)"
    ws["G11"].number_format = "#,##0.00"
    ws["G11"].font = Font(bold=True)
    ws.merge_cells("A13:H13")
    ws["A13"] = "Payment: Irrevocable Letter of Credit at sight"
    ws.merge_cells("A14:H14")
    ws["A14"] = "Quality: ISO 15243 test report enclosed with each lot (report ref PMC-TR-15243-2026)"
    ws.merge_cells("A15:H15")
    ws["A15"] = "Validity: 30 days. E&OE."
    for col, w in zip("ABCDEFGH", [6, 15, 26, 13, 7, 16, 14, 10]):
        ws.column_dimensions[col].width = w
    p = OUT / "vendor_D_PacificMotion_quote.xlsx"
    wb.save(p)
    return p


# --------------------------------------------------------------------------- Vendor E
def vendor_e() -> Path:
    """Plain email body, no attachment. Terse, 'same terms as our last order'."""
    body = ("Hi,\n\n"
            "Can supply all three types - std Rs 2300, heavy 4700, high speed 6350 per pc, dispatch in 5 weeks. "
            "Same terms as our last order.\n\n"
            "Rakesh\nVardhman Industrial Supplies\n")
    p = OUT / "vendor_E_Vardhman_email_body.txt"
    p.write_text(body)
    return p


# --------------------------------------------------------------------------- Vendor F
def vendor_f() -> Path:
    """Generic printed rate card, not addressed to this RFQ. Cheapest, but 12-13 week lead time.
    30 days from GRN. Ex-godown Chennai."""
    p = OUT / "vendor_F_SriLakshmi_ratecard.pdf"
    doc = SimpleDocTemplate(str(p), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm)
    s = lambda size, font="DejaVu", color=colors.black: ParagraphStyle(  # noqa: E731
        "x", fontName=font, fontSize=size, textColor=color, leading=size * 1.35)
    story = [
        Paragraph("SRI LAKSHMI BEARING HOUSE", s(22, "DejaVu-Bold", colors.HexColor("#8B0000"))),
        Paragraph("Authorised Stockist & Assembler · 14, Armenian Street, Parrys, Chennai 600001 · "
                  "Ph 044-0000 0000 · slbh.chennai@mail.example", s(8.5)),
        Spacer(1, 5 * mm),
        Paragraph("RATE CARD — BEARING ASSEMBLIES (Effective 1 September 2026)", s(13, "DejaVu-Bold")),
        Spacer(1, 4 * mm),
    ]
    data = [["SKU", "Description", "Series", "Rate ₹ / pc"],
            ["SL-A6205", "Light duty DGBB assembly", "6205", "650"],
            ["SL-A6308", "Medium duty DGBB assembly", "6308", "1,180"],
            ["SL-A6212", "General duty DGBB assembly with housing", "6212", "2,150"],
            ["SL-A51110", "Thrust ball bearing assembly", "51110", "890"],
            ["SL-ANU316", "Heavy duty CRB assembly with housing", "NU316", "4,300"],
            ["SL-A7014P4", "High speed ACBB assembly, P4 precision", "7014", "5,900"],
            ["SL-A22220", "Spherical roller bearing assembly", "22220", "7,800"],
            ["SL-APB208", "Pillow block unit", "UCP208", "1,420"]]
    tbl = Table(data, colWidths=[28 * mm, 82 * mm, 24 * mm, 30 * mm])
    tbl.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "DejaVu", 10),
        ("FONT", (0, 0), (-1, 0), "DejaVu-Bold", 10),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4D6D6")),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("ALIGN", (3, 0), (3, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [tbl, Spacer(1, 6 * mm)]
    for line in [
        "• Standard lead time for assemblies: 12–13 weeks from PO. Loose bearings ex-stock.",
        "• Payment: 30 days from GRN.",
        "• Prices ex-godown Chennai. Packing & forwarding 1%, GST 18% extra. Freight to pay.",
        "• Certifications: ISO 9001:2015 (cert IQC/9001/7741, valid to 30-Jun-2028); "
        "assemblies conform to ISO 15243 (cert SLBH-15243-221, valid to 31-Dec-2027).",
        "• Rates subject to change without notice. Minimum order value ₹25,000.",
    ]:
        story.append(Paragraph(line, s(10)))
    doc.build(story)
    return p


def _perspective_coeffs(src, dst):
    """Coefficients for PIL's PERSPECTIVE transform mapping output (dst) -> input (src)."""
    m = []
    for (x, y), (u, v) in zip(dst, src):
        m.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        m.append([0, 0, 0, x, y, 1, -v * x, -v * y])
    a = np.array(m, dtype=float)
    b = np.array(src, dtype=float).reshape(8)
    return np.linalg.solve(a, b).tolist()


def vendor_f_simulated_photo(pdf: Path) -> Path:
    """Stand-in for a phone photo of the printed rate card: angled, uneven light, soft focus, JPEG.
    Replace with a real photograph for the demo; the pipeline treats both the same."""
    import pypdfium2 as pdfium
    page = pdfium.PdfDocument(str(pdf))[0]
    img = page.render(scale=2.0).to_pil().convert("RGB")
    w, h = img.size
    canvas = Image.new("RGB", (int(w * 1.25), int(h * 1.2)), (92, 78, 64))  # wooden desk
    canvas.paste(img, (int(w * 0.12), int(h * 0.1)))
    cw, ch = canvas.size
    src = [(0, 0), (cw, 0), (cw, ch), (0, ch)]
    dst = [(cw * 0.06, ch * 0.03), (cw * 0.97, ch * 0.09), (cw * 0.90, ch * 0.98), (cw * 0.01, ch * 0.92)]
    warped = canvas.transform((cw, ch), Image.PERSPECTIVE, _perspective_coeffs(src, dst), Image.BICUBIC,
                              fillcolor=(92, 78, 64))
    warped = warped.rotate(-2.5, resample=Image.BICUBIC, fillcolor=(92, 78, 64))
    grad = np.linspace(1.05, 0.72, cw)[None, :, None] * np.linspace(1.0, 0.85, ch)[:, None, None]
    arr = np.clip(np.asarray(warped, dtype=float) * grad * np.array([1.0, 0.97, 0.9]), 0, 255)
    rng = np.random.default_rng(7)
    arr = np.clip(arr + rng.normal(0, 6, arr.shape), 0, 255).astype(np.uint8)
    photo = Image.fromarray(arr).filter(ImageFilter.GaussianBlur(1.1))
    photo = ImageEnhance.Contrast(photo).enhance(0.9).resize((int(cw * 0.6), int(ch * 0.6)))
    p = OUT / "vendor_F_SriLakshmi_ratecard_photo_SIMULATED.jpg"
    photo.save(p, quality=72)
    return p


if __name__ == "__main__":
    for f in (vendor_a, vendor_b, vendor_c, vendor_d, vendor_e):
        print(f())
    pdf = vendor_f()
    print(pdf)
    print(vendor_f_simulated_photo(pdf))
