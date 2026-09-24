"""Generate six vendor replies for an arbitrary RFQ, each vendor with a consistent character:

A  Shreeji (xlsx)      reliable incumbent: quotes everything, INR, delivered, 45 days, certs met, a bit dearer
B  Nordlager (PDF)     EUR, ex-works, 2/10 net 30 in small print, near-equivalent (DIN/EN) standards
C  Kaveri (docx)       prose letter, declines one item, 50% advance, free freight above ₹5 lakh,
                       one certificate expiring just before need-by
D  Pacific (xlsx)      messy sheet, USD, small parts priced per pack of 2, LC at sight, no freight, no ISO 9001
E  Vardhman (email)    two sentences, prices only, "same terms as our last order"
F  Sri Lakshmi (photo) generic rate card, cheapest, but too slow (12+ weeks)

Lead times are set against the time actually available before each need-by date: A, C and E can
deliver in time, B, D and F cannot — so every RFQ shows a mix of vendors who can and can't deliver.
"""
import hashlib
import random
from datetime import date, timedelta
from email.message import EmailMessage
from pathlib import Path

import numpy as np
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageEnhance, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

EUR, USD = 98.10, 88.40
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Indicative market price per unit, INR — the simulated vendors' starting point.
BASE_PRICE = {
    ("MAT-001", "—"): 38000, ("MAT-002", "—"): 14500, ("MAT-003", "—"): 26000, ("MAT-004", "—"): 9200,
    ("MAT-005", "Standard"): 2400, ("MAT-005", "Heavy-duty"): 4800, ("MAT-005", "High-speed"): 6500,
    ("MAT-006", "Single cartridge"): 18500, ("MAT-006", "Double cartridge"): 34000,
    ("MAT-007", "200 kW"): 610000, ("MAT-007", "110 kW"): 380000, ("MAT-007", "75 kW"): 265000,
    ("MAT-008", "—"): 12500, ("MAT-009", "—"): 420000, ("MAT-010", "—"): 3800,
}
# Typical manufacturing lead time, weeks
TYPICAL_WEEKS = {"MAT-001": 8, "MAT-002": 7, "MAT-003": 6, "MAT-004": 5, "MAT-005": 5, "MAT-006": 6,
                 "MAT-007": 10, "MAT-008": 4, "MAT-009": 9, "MAT-010": 2}
# What a European supplier certifies to instead of the Indian/ISO standard asked for
NEAR_EQUIV = {"ISO 15243": "DIN 628", "IS 210 Gr FG260": "EN-GJL-250 (EN 1561)", "IS 2062 E350 BR": "EN 10025 S355JR",
              "IS 2062": "EN 10025", "API 682 4th ed.": "EN 12756", "IS 12615": "IEC 60034-1",
              "IS 305": "DIN EN 1982 CuAl10Fe5Ni5", "IS 1367 property class 8.8": "DIN EN ISO 898-1 8.8"}

FONT_DIR = Path(__file__).resolve().parent / "fonts"  # bundled: hosts (e.g. Streamlit Cloud) lack DejaVu
FONT, FONT_B = "Helvetica", "Helvetica-Bold"  # replaced by DejaVu (which has the ₹ glyph) once registered
RUPEE = "Rs."
_fonts_ready = False


def _fonts():
    """Use the bundled DejaVu fonts; if they cannot be loaded, fall back to built-in Helvetica and 'Rs.'."""
    global _fonts_ready, FONT, FONT_B, RUPEE
    if _fonts_ready:
        return
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        FONT, FONT_B, RUPEE = "DejaVu", "DejaVu-Bold", "₹"
    except Exception:  # noqa: BLE001
        pass
    _fonts_ready = True


def _rng(*parts) -> random.Random:
    return random.Random(int(hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest()[:12], 16))


def _label(it: dict) -> str:
    return it["material"] + ("" if it["variant"] == "—" else f" — {it['variant']}")


def _standards(it: dict) -> list[str]:
    return [p.strip() for p in it["quality_standard"].split("+") if p.strip()]


def _is_supplier_qms(std: str) -> bool:
    s = std.lower()
    return "9001" in s or "ce marking" in s


ON_TIME = {"A", "C", "E"}  # these vendors fit the need-by date; B, D and F do not


def _weeks_for(code: str, vendor: str, window_days: int) -> tuple[int, int]:
    """(lead weeks lo, hi) as the vendor would state it, relative to the time available before need-by.
    A, C and E always fit (ex-stock when time is short); B, D and F always miss — so every RFQ shows a mix."""
    if vendor in ON_TIME:
        w = int(max(window_days, 0) * {"A": 0.75, "C": 0.6, "E": 0.85}[vendor] // 7)
        return w, w
    w = -(-max(window_days, 0) * {"B": 1.25, "D": 1.15, "F": 1.6}[vendor] // 7) + 2
    w = max(int(w), TYPICAL_WEEKS.get(code, 6))
    return (max(12, w), max(13, w + 1)) if vendor == "F" else (w, w)


def _lead_text(lo: int, hi: int, style: str) -> str:
    if hi == 0:
        return {"xlsx": "Ex-stock", "prose": "immediately from stock", "short": "ex-stock"}[style]
    if lo != hi:
        return f"{lo}–{hi} weeks"
    return {"xlsx": f"{lo} weeks", "prose": f"{lo * 7} days", "short": f"{lo} wks"}[style]


def _price(it: dict, vendor: str, rfq_id: str) -> float:
    mult = {"A": 1.05, "B": 1.00, "C": 0.965, "D": 0.99, "E": 0.95, "F": 0.88}[vendor]
    noise = _rng(rfq_id, vendor, it["material_code"], it["variant"]).uniform(-0.02, 0.02)
    return BASE_PRICE.get((it["material_code"], it["variant"]), 10000) * (mult + noise)


def _de(x: float) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# --------------------------------------------------------------------------- A · xlsx
def _vendor_a(rfq, items, po, out: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Offer"
    ws["A1"] = "SHREEJI BEARINGS & TRANSMISSION PVT LTD"
    ws["A1"].font = Font(bold=True, size=14, color="7B2C2C")
    ws["A2"] = "Survey No. 118, Metoda GIDC, Rajkot 360021"
    ws["A4"] = f"Offer No: SBT/OF/26-27/{_rng(rfq['rfq_id'], 'A').randint(600, 999)}"
    ws["E4"] = f"Your Ref: {rfq['rfq_id']}"
    hdr = ["S.No", "Item Code", "Description", "Offered Qty", "UOM", "Rate (INR)", "Lead Time", "Remarks"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=6, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7B2C2C")
    r = 7
    for n, it in enumerate(reversed(items), 1):  # their own order
        lo, hi = _weeks_for(it["material_code"], "A", (date.fromisoformat(it["need_by"]) - po).days)
        code = f"SBT-{it['material_code'][-3:]}-{(it['variant'][:2] if it['variant'] != '—' else 'GN').upper()}"
        desc = it["material"].lower() + ("" if it["variant"] == "—" else f", {it['variant'].lower()} type")
        vals = [n, code, desc, it["qty"], it["uom"].capitalize(), round(_price(it, "A", rfq["rfq_id"]), 0),
                _lead_text(lo, hi, "xlsx"), ""]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v).border = BOX
        r += 1
    r += 1
    certs = sorted({s for it in items for s in _standards(it)})
    for label, val in [("Price basis", "FOR your Chakan plant — freight & insurance included. GST extra."),
                       ("Payment", "45 days net"), ("Validity", "60 days"),
                       ("Certification", "; ".join(f"{s}: cert SBT/QC/{abs(hash(s)) % 9000 + 1000}, valid to "
                                                   "31-Mar-2028" for s in certs))]:
        ws.cell(row=r, column=2, value=label).font = Font(bold=True)
        ws.cell(row=r, column=3, value=val)
        r += 1
    for col, w in zip("ABCDEFGH", [6, 16, 48, 12, 7, 12, 12, 30]):
        ws.column_dimensions[col].width = w
    p = out / f"vendor_A_Shreeji_offer_{rfq['rfq_id']}.xlsx"
    wb.save(p)
    return p


# --------------------------------------------------------------------------- B · PDF, EUR
def _vendor_b(rfq, items, po, out: Path) -> Path:
    _fonts()
    p = out / f"vendor_B_Nordlager_Angebot_{rfq['rfq_id']}.pdf"
    doc = SimpleDocTemplate(str(p), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm,
                            bottomMargin=15 * mm)
    s = lambda size, font=None, color=colors.black, lead=None: ParagraphStyle(  # noqa: E731
        "x", fontName=font or FONT, fontSize=size, textColor=color, leading=lead or size * 1.3)
    story = [Paragraph("NORDLAGER", s(24, FONT_B, colors.HexColor("#0B4F8A"))),
             Paragraph("Antriebstechnik GmbH · Heilbronner Straße 211 · 70191 Stuttgart", s(8, color=colors.grey)),
             Spacer(1, 8 * mm),
             Paragraph(f"<b>Angebot / Quotation NL-Q-2026-{_rng(rfq['rfq_id'], 'B').randint(3000, 3999)}</b> — "
                       f"your enquiry {rfq['rfq_id']}", s(12, FONT_B)),
             Spacer(1, 4 * mm)]
    data = [["Pos.", "Description", "Qty", "Unit price EUR", "Total EUR"]]
    for n, it in enumerate(items, 1):
        eur = round(_price(it, "B", rfq["rfq_id"]) / EUR, 2)
        data.append([str(n * 10), Paragraph(_label(it), s(8.5)), str(it["qty"]), _de(eur), _de(eur * it["qty"])])
    t = Table(data, colWidths=[12 * mm, 90 * mm, 14 * mm, 28 * mm, 28 * mm])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), FONT, 8.5), ("FONT", (0, 0), (-1, 0), FONT_B, 8.5),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F2")),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("ALIGN", (2, 0), (-1, -1), "RIGHT")]))
    weeks = max(_weeks_for(it["material_code"], "B", (date.fromisoformat(it["need_by"]) - po).days)[1]
                for it in items)
    stds = sorted({s_ for it in items for s_ in _standards(it)})
    cited = [NEAR_EQUIV.get(x, x) for x in stds if not _is_supplier_qms(x)]
    story += [t, Spacer(1, 5 * mm),
              Paragraph(f"<b>Delivery:</b> {weeks} weeks after receipt of order, ex works.", s(10)),
              Paragraph("<b>Quality:</b> manufactured and tested in accordance with " + ", ".join(cited)
                        + ". Company certified ISO 9001:2015 (DQS 074512 QM15, valid until 31.05.2028).", s(10)),
              Spacer(1, 40 * mm),
              Paragraph("Terms of payment: 2% discount if paid within 10 days of invoice date, otherwise net 30 days "
                        "from invoice date. Prices EXW Stuttgart (Incoterms® 2020); freight, insurance and duties "
                        "for buyer's account. Offer valid 45 days.", s(5.5, color=colors.HexColor("#808080"),
                                                                           lead=7))]
    doc.build(story)
    return p


# --------------------------------------------------------------------------- C · docx prose
def _vendor_c(rfq, items, po, out: Path) -> Path:
    d = Document()
    d.styles["Normal"].font.size = Pt(11)
    h = d.add_paragraph().add_run("KAVERI PRECISION COMPONENTS")
    h.bold, h.font.size = True, Pt(16)
    d.add_paragraph(f"Ref: KPC/MKT/2026/{_rng(rfq['rfq_id'], 'C').randint(1000, 1999)} — Sub: your enquiry "
                    f"{rfq['rfq_id']}")
    d.add_paragraph("Dear Sir,")
    d.add_paragraph("We thank you for your enquiry and are pleased to submit our offer as below.")
    declined = items[-1] if len(items) > 1 else None
    for it in items:
        if it is declined:
            continue
        lo, hi = _weeks_for(it["material_code"], "C", (date.fromisoformat(it["need_by"]) - po).days)
        d.add_paragraph(f"For the {_label(it).lower()}, we offer ₹{_price(it, 'C', rfq['rfq_id']):,.0f} per "
                        f"{it['uom'].rstrip('s')} for the full quantity of {it['qty']}, deliverable "
                        f"{_lead_text(lo, hi, 'prose')} from your purchase order.")
    if declined:
        d.add_paragraph(f"We regret that we do not presently manufacture the {_label(declined).lower()} and are "
                        "unable to quote for it.")
    d.add_paragraph("As regards commercial terms, we request 50% advance along with the purchase order, with the "
                    "balance payable on delivery. Freight will be free to your Chakan plant for any single despatch "
                    "above ₹5 lakh in value; below this, freight at actuals.")
    first = next(it for it in items if it is not declined)
    exp = date.fromisoformat(first["need_by"]) - timedelta(days=6)
    std = _standards(first)[0]
    d.add_paragraph(f"Our products conform to {', '.join(sorted({s for it in items for s in _standards(it)}))}. "
                    f"Our {std} certificate (KPC/{std.replace(' ', '')}/0457, Bureau Veritas) is valid until "
                    f"{exp:%d %B %Y}; all other certificates are valid till 31 August 2027.")
    d.add_paragraph("Yours faithfully,\nS. Muthukumar, Manager — Marketing")
    p = out / f"vendor_C_Kaveri_letter_{rfq['rfq_id']}.docx"
    d.save(p)
    return p


# --------------------------------------------------------------------------- D · messy xlsx, USD, packs
def _vendor_d(rfq, items, po, out: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:H1")
    ws["A1"] = "PACIFIC MOTION COMPONENTS PTE LTD"
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells("A2:H2")
    ws["A2"] = "21 Tuas South Ave 3, Singapore 637419"
    ws.merge_cells("A3:H3")
    ws["A3"] = "Q U O T A T I O N"
    ws["A4"] = f"Quote #: PMC-{_rng(rfq['rfq_id'], 'D').randint(1000, 9999)}"
    ws.merge_cells("A5:H5")
    ws["A5"] = "To: Deccan Flow Systems (India)  Attn: Procurement"
    hdr = ["Item", "Part No.", "Desc.", "UOM", "Qty", "Unit Price (USD)", "Amount (USD)", "Lead Time"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=7, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    r = 8
    for n, it in enumerate(items, 1):
        unit_inr = _price(it, "D", rfq["rfq_id"])
        pack = 2 if unit_inr < 20000 and it["qty"] % 2 == 0 else 1
        lo, hi = _weeks_for(it["material_code"], "D", (date.fromisoformat(it["need_by"]) - po).days)
        vals = [n, f"PMC-{it['material_code'][-3:]}{n}", _label(it).lower(),
                "PACK (2 PCS)" if pack == 2 else "PC", it["qty"] // pack, round(unit_inr * pack / USD, 2)]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v).border = BOX
        ws.cell(row=r, column=7, value=f"=E{r}*F{r}")
        ws.cell(row=r, column=8, value=_lead_text(lo, hi, "short"))
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.cell(row=r, column=1, value="Payment: Irrevocable Letter of Credit at sight")
    prod = sorted({s for it in items for s in _standards(it) if not _is_supplier_qms(s)})
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=8)
    ws.cell(row=r + 1, column=1, value="Quality: test reports to " + ", ".join(prod) + " enclosed with each lot")
    p = out / f"vendor_D_PacificMotion_quote_{rfq['rfq_id']}.xlsx"
    wb.save(p)
    return p


# --------------------------------------------------------------------------- E · terse email
def _vendor_e(rfq, items, po) -> str:
    parts = ", ".join(f"{_label(it).lower()} Rs {_price(it, 'E', rfq['rfq_id']):,.0f}" for it in items)
    weeks = min(_weeks_for(it["material_code"], "E", (date.fromisoformat(it["need_by"]) - po).days)[1]
                for it in items)
    when = "ex-stock" if weeks == 0 else f"dispatch in {weeks} weeks"
    return (f"Hi,\n\nCan supply - {parts} per pc, {when}. Same terms as our last order.\n\n"
            "Rakesh\nVardhman Industrial Supplies\n")


# --------------------------------------------------------------------------- F · rate card → photo
def _vendor_f(rfq, items, po, out: Path) -> Path:
    _fonts()
    pdf = out / f"vendor_F_ratecard_{rfq['rfq_id']}.pdf"
    doc = SimpleDocTemplate(str(pdf), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=14 * mm)
    s = lambda size, font=None: ParagraphStyle("x", fontName=font or FONT, fontSize=size, leading=size * 1.35)  # noqa
    rows = [["SKU", "Description", f"Rate {RUPEE} / pc"]]
    distract = [("SL-X6205", "Light duty DGBB assembly", 650), ("SL-X51110", "Thrust ball bearing assembly", 890)]
    for it in items:
        rows.append([f"SL-{it['material_code'][-3:]}{(it['variant'][:1] if it['variant'] != '—' else 'G')}",
                     _label(it), f"{_price(it, 'F', rfq['rfq_id']):,.0f}"])
    rows += [[a, b, f"{c:,}"] for a, b, c in distract]
    t = Table(rows, colWidths=[30 * mm, 100 * mm, 34 * mm])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), FONT, 10), ("FONT", (0, 0), (-1, 0), FONT_B, 10),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4D6D6")),
                           ("GRID", (0, 0), (-1, -1), 0.6, colors.black), ("ALIGN", (2, 0), (2, -1), "RIGHT"),
                           ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story = [Paragraph("SRI LAKSHMI BEARING HOUSE", s(20, FONT_B)),
             Paragraph("Authorised Stockist & Assembler · Parrys, Chennai 600001", s(8.5)), Spacer(1, 4 * mm),
             Paragraph("RATE CARD (Effective 1 September 2026)", s(13, FONT_B)), Spacer(1, 3 * mm), t,
             Spacer(1, 5 * mm)]
    lo, hi = max((_weeks_for(it["material_code"], "F", (date.fromisoformat(it["need_by"]) - po).days)
                  for it in items), key=lambda x: x[1])
    for line in [f"• Standard lead time: {lo}–{hi} weeks from PO, all items.", "• Payment: 30 days from GRN.",
                 "• Prices ex-godown Chennai. GST 18% extra. Freight to pay.",
                 "• Certifications: " + ", ".join(sorted({x for it in items for x in _standards(it)}))
                 + " — all valid to 31-Dec-2027."]:
        story.append(Paragraph(line, s(10)))
    doc.build(story)
    return _photograph(pdf, out / f"vendor_F_ratecard_photo_{rfq['rfq_id']}.jpg")


def _coeffs(src, dst):
    m = []
    for (x, y), (u, v) in zip(dst, src):
        m.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        m.append([0, 0, 0, x, y, 1, -v * x, -v * y])
    return np.linalg.solve(np.array(m, float), np.array(src, float).reshape(8)).tolist()


def _photograph(pdf: Path, dest: Path) -> Path:
    import pypdfium2 as pdfium
    img = pdfium.PdfDocument(str(pdf))[0].render(scale=2.0).to_pil().convert("RGB")
    w, h = img.size
    canvas = Image.new("RGB", (int(w * 1.25), int(h * 1.2)), (92, 78, 64))
    canvas.paste(img, (int(w * 0.12), int(h * 0.1)))
    cw, ch = canvas.size
    src = [(0, 0), (cw, 0), (cw, ch), (0, ch)]
    dst = [(cw * 0.06, ch * 0.03), (cw * 0.97, ch * 0.09), (cw * 0.90, ch * 0.98), (cw * 0.01, ch * 0.92)]
    warped = canvas.transform((cw, ch), Image.PERSPECTIVE, _coeffs(src, dst), Image.BICUBIC, fillcolor=(92, 78, 64))
    warped = warped.rotate(-2.5, resample=Image.BICUBIC, fillcolor=(92, 78, 64))
    grad = np.linspace(1.05, 0.72, cw)[None, :, None] * np.linspace(1.0, 0.85, ch)[:, None, None]
    arr = np.clip(np.asarray(warped, float) * grad * np.array([1.0, 0.97, 0.9]), 0, 255)
    arr = np.clip(arr + np.random.default_rng(7).normal(0, 6, arr.shape), 0, 255).astype(np.uint8)
    photo = Image.fromarray(arr).filter(ImageFilter.GaussianBlur(1.0))
    photo = ImageEnhance.Contrast(photo).enhance(0.9).resize((int(cw * 0.6), int(ch * 0.6)))
    photo.save(dest, quality=74)
    return dest


# --------------------------------------------------------------------------- emails
def _eml(dest: Path, frm, to, subject, body, attachment: Path | None = None) -> Path:
    import mimetypes
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = frm, to, subject
    m["Date"] = "Wed, 30 Sep 2026 11:00:00 +0530"
    m["Message-ID"] = f"<{hashlib.md5(str(dest).encode()).hexdigest()}@sim.example>"
    m.set_content(body)
    if attachment:
        mt, st = (mimetypes.guess_type(attachment.name)[0] or "application/octet-stream").split("/", 1)
        m.add_attachment(attachment.read_bytes(), maintype=mt, subtype=st, filename=attachment.name)
    dest.write_bytes(m.as_bytes())
    return dest


def simulate_replies(rfq: dict, vendor_codes: list[str], reply_to: dict[str, str], po: date,
                     inbox: Path, work: Path) -> list[Path]:
    """Write one reply email per vendor into `inbox`. reply_to: vendor code → correlation reply address."""
    work.mkdir(parents=True, exist_ok=True)
    items, rid = rfq["lines"], rfq["rfq_id"]
    subj = f"Re: {rid} — Request for quotation: {rfq['title']}"
    out = []
    for v in vendor_codes:
        dest = inbox / f"sim_{rid}_{v}.eml"
        if dest.exists():
            continue
        if v == "A":
            out.append(_eml(dest, "Shreeji Sales <sales@shreeji-bearings.example>", reply_to[v], subj,
                            "Dear Sir,\n\nPlease find our offer attached.\n\nRegards,\nNilesh Patel",
                            _vendor_a(rfq, items, po, work)))
        elif v == "B":
            out.append(_eml(dest, "Jürgen Albrecht <j.albrecht@nordlager.example>", "procurement@buyer.example",
                            f"Angebot zu Ihrer Anfrage {rid}", "Dear Sir or Madam,\n\nplease find attached our "
                            "quotation.\n\nKind regards\nJürgen Albrecht", _vendor_b(rfq, items, po, work)))
        elif v == "C":
            out.append(_eml(dest, "Kaveri Precision <info@kaveriprecision.example>", reply_to[v], subj,
                            "Sir,\n\nKindly find our offer letter attached.\n\nS. Muthukumar",
                            _vendor_c(rfq, items, po, work)))
        elif v == "D":
            out.append(_eml(dest, "PMC Sales Asia <sales.asia@pacificmotion.example>", "procurement@buyer.example",
                            "Quotation - as requested", "Hi,\n\nAttached our best quote.\n\nMelvin Tan",
                            _vendor_d(rfq, items, po, work)))
        elif v == "E":
            out.append(_eml(dest, "Rakesh Jain <rakesh@vardhman-ind.example>", reply_to[v], subj,
                            _vendor_e(rfq, items, po)))
        elif v == "F":
            out.append(_eml(dest, "Lakshmi Bearings <lakshmibearings.chn@webmail.example>",
                            "procurement@buyer.example", "rate card", "Sir, pls find our latest rate card. Regards",
                            _vendor_f(rfq, items, po, work)))
    return out
