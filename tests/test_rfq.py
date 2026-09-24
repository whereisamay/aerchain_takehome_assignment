import sys
from datetime import date
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

from openpyxl import load_workbook  # noqa: E402

from demand import compute_demand, po_date_for  # noqa: E402
from master_data import MATERIALS  # noqa: E402
from rfq import build_rfq, validate_rfq  # noqa: E402
from rfq_xlsx import render_rfq_xlsx  # noqa: E402

TARGET = date(2027, 2, 10)
DEMAND = compute_demand({"PS-200": 12, "PS-075": 20}, TARGET)
PO = po_date_for(TARGET)
TODAY = date(2026, 9, 23)


def test_single_material_rfq_keeps_material_number():
    rfq = build_rfq(DEMAND, ["MAT-005"], PO, set(), today=TODAY)
    assert rfq["rfq_id"] == "RFQ-2026-MCH-005"
    assert rfq["title"] == "Bearing assembly"
    assert rfq["close_date"] == "2026-10-02"
    assert {l["variant"]: l["qty"] for l in rfq["lines"]} == {"Standard": 40, "Heavy-duty": 24, "High-speed": 20}
    assert validate_rfq(rfq) == []


def test_all_p1_parts_on_one_rfq():
    p1 = [c for c, m in MATERIALS.items() if m.phase == "P1"]
    rfq = build_rfq(DEMAND, p1, PO, set(), today=TODAY)
    assert rfq["rfq_id"] == "RFQ-2026-MCH-101"
    assert rfq["title"].startswith("P1 · Casting and machining — 4 materials")
    assert len(rfq["lines"]) == 4 and all(l["need_by"] == "2026-10-07" for l in rfq["lines"])
    assert validate_rfq(rfq) == []
    # next multi-material RFQ takes the next free number
    assert build_rfq(DEMAND, p1, PO, {"RFQ-2026-MCH-101"}, today=TODAY)["rfq_id"] == "RFQ-2026-MCH-102"


def test_everything_generates_and_exports():
    rfq = build_rfq(DEMAND, list(MATERIALS), PO, set(), today=TODAY)
    assert len(rfq["lines"]) == 14
    wb = load_workbook(BytesIO(render_rfq_xlsx(rfq)))
    assert wb.sheetnames == ["Terms", "Line items", "Certifications"]
    ws = wb["Line items"]
    assert ws.cell(row=6, column=2).value == "Material"
    assert ws.cell(row=7, column=2).value == "Pump casing, cast iron"
    assert ws.cell(row=7, column=8).fill.fgColor.rgb.endswith("FFF2CC")  # unit price is a shaded input
    certs = wb["Certifications"]
    assert {certs.cell(row=r, column=1).value for r in range(7, 30)} >= {"Bearing assembly", "Mechanical seal"}


def test_validation_catches_bad_edits():
    rfq = build_rfq(DEMAND, ["MAT-005", "MAT-001"], PO, set(), today=TODAY)
    rfq["lines"][0]["variant"] = "Ceramic"
    rfq["lines"][1]["qty"] = 0
    rfq["lines"].append(dict(rfq["lines"][2]))
    problems = validate_rfq(rfq)
    assert any("Ceramic" in p for p in problems)
    assert any("positive" in p for p in problems)
    assert any("more than one line" in p for p in problems)
